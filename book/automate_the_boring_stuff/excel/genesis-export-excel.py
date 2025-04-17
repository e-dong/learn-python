import requests
import os
from pathlib import Path
import openpyxl
from progress_bar import ProgressBar

bible_id = 'de4e12af7f28f599-02' # KJV-2
with open(Path(os.getenv("HOME")) / ".api.bible.creds") as f:
    apikey = f.read()

def generate_sheet(wb, chapter):
    idx, chapter_data = chapter
    if chapter_data['number'] == "intro":
       return 
    sheet = wb.create_sheet(index=idx, title=chapter_data['reference'])
    verse_response = requests.get(f"https://api.scripture.api.bible/v1/bibles/{bible_id}/chapters/{chapter_data['id']}/verses", headers={ "api-key": apikey })
    verses = verse_response.json()['data']
    for verse_idx, verse in enumerate(verses):
        verse_content_response = requests.get(f"https://api.scripture.api.bible/v1/bibles/{bible_id}/verses/{verse['id']}?content-type=text", headers={ "api-key": apikey })
        verse_obj = verse_content_response.json()['data']
        sheet.cell(row=verse_idx + 1, column=1, value=verse_obj['content'])


def main():
    chapter_response = requests.get(f"https://api.scripture.api.bible/v1/bibles/{bible_id}/books/GEN/chapters", headers={ "api-key": apikey })
    gen_chapters = chapter_response.json()["data"]

    pb = ProgressBar(len(gen_chapters))

    # new workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for chapter in enumerate(gen_chapters):
        generate_sheet(wb, chapter)
        pb.render()
    wb.save('Genesis.xlsx')

if __name__ == "__main__":
    main()
