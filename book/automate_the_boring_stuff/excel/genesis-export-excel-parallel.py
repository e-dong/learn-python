import os
import sys
import traceback
from concurrent.futures import ALL_COMPLETED, ThreadPoolExecutor, as_completed, wait
from pathlib import Path

import openpyxl
import requests
from progress_bar import ProgressBar

bible_id = "de4e12af7f28f599-02"  # KJV-2
with open(Path(os.getenv("HOME")) / ".api.bible.creds") as f:
    apikey = f.read()


def generate_sheet(wb, chapter):
    idx, chapter_data = chapter
    if chapter_data["number"] == "intro":
        return
    sheet = wb.create_sheet(index=idx, title=chapter_data["reference"])
    verse_response = requests.get(
        f"https://api.scripture.api.bible/v1/bibles/{bible_id}/chapters/{chapter_data['id']}/verses",
        headers={"api-key": apikey},
    )
    verse_response.raise_for_status()
    verses = verse_response.json()["data"]
    for verse_idx, verse in enumerate(verses):
        verse_content_response = requests.get(
            f"https://api.scripture.api.bible/v1/bibles/{bible_id}/verses/{verse['id']}?content-type=text",
            headers={"api-key": apikey},
        )
        verse_content_response.raise_for_status()
        verse_obj = verse_content_response.json()["data"]
        sheet.cell(row=verse_idx + 1, column=1, value=verse_obj["content"])


def main():
    print("Fetch all chapters in book...")
    chapter_response = requests.get(
        f"https://api.scripture.api.bible/v1/bibles/{bible_id}/books/GEN/chapters",
        headers={"api-key": apikey},
    )
    chapter_response.raise_for_status()
    gen_chapters = chapter_response.json()["data"]

    pb = ProgressBar(len(gen_chapters))

    # new workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    with ThreadPoolExecutor(max_workers=10) as exec:
        futures = [
            exec.submit(generate_sheet, wb, chapter)
            for chapter in enumerate(gen_chapters)
        ]
        for future in as_completed(futures):
            try:
                future.result()
            except Exception as e:
                trace = traceback.format_exc()
                print("%s\n%s" % (e, trace))
                sys.exit(1)
            pb.render()
        wait(futures, return_when=ALL_COMPLETED)
        print("Saving excel file!")
        wb.save("Genesis.xlsx")
        wb.close()


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        trace = traceback.format_exc()
        print("%s\n%s" % (e, trace))
